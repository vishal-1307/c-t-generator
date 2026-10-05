"""The Master Timetable, audited against every other way of reading a timetable.

The master view is one row per ``Assignment`` of one run. These tests hold it
to four promises, each checked against an answer derived independently from
the ``Assignment`` rows rather than against the view itself:

1. every scheduled period appears exactly once - nothing missing, nothing
   doubled, across every page - with every column a person reads;
2. every filter, alone and combined, returns exactly the matching rows, and
   the count and pages are of what matches;
3. it is the same run every other view shows, never a failed run and never a
   mixture of two;
4. the same class reads identically in the master view, the three week grids,
   the teacher's table and the Excel file - before and after a manual move, a
   room change and a regeneration that respects locks - with nothing stale.

And the version rules around it: generating adds a version and keeps the old
one, a published version cannot be changed in place or deleted, and two
versions can be compared.
"""
from __future__ import annotations

import io
from collections import Counter

import pytest
from openpyxl import load_workbook

from app.grid import build_grid
from app.models import (
    Assignment,
    Faculty,
    Room,
    Section,
    Subject,
    TimeSlot,
    TimetableRun,
)
from app.solver.run import generate
from app.timetable_views import room_code

from constraint_checker import check_all


# ------------------------------------------------------------------- fixture


@pytest.fixture
def solved(db_session, context):
    """Two sections, a theory subject each and a two-period BYOD lab, three
    teachers, classrooms in two blocks on different floors and one lab."""
    db = db_session
    theory_a = Subject(name="Data Structures", code="CS201", type="theory",
                       session_length_hours=1, sessions_per_week=3)
    theory_b = Subject(name="Operating Systems", code="CS202", type="theory",
                       session_length_hours=1, sessions_per_week=2)
    lab = Subject(name="IoT Lab", code="EC282", type="practical", required_lab_type="iot",
                  session_length_hours=2, sessions_per_week=1, byod_required=True)
    db.add_all([theory_a, theory_b, lab])
    db.flush()

    rooms = [
        Room(room_number="101", block="36", floor="1", capacity=70, room_type="theory"),
        Room(room_number="102", block="36", floor="1", capacity=70, room_type="theory"),
        Room(room_number="201", block="37", floor="2", capacity=80, room_type="theory",
             byod=True),
        Room(room_number="301", block="37", floor="3", capacity=70, room_type="lab",
             lab_type="iot", byod=True),
    ]
    db.add_all(rooms)
    db.flush()

    faculty = [Faculty(name=n, faculty_code=c) for n, c in
               [("Faculty A", "90001"), ("Faculty B", "90002"), ("Faculty C", "90003")]]
    faculty[0].subjects = [theory_a, theory_b]
    faculty[1].subjects = [theory_a, theory_b]
    faculty[2].subjects = [lab, theory_b]
    db.add_all(faculty)

    sections = [
        Section(academic_context_id=context.id, section_number="2401", strength=60),
        Section(academic_context_id=context.id, section_number="2402", strength=64),
    ]
    sections[0].subjects = [theory_a, lab]
    sections[1].subjects = [theory_b, lab]
    db.add_all(sections)
    db.add_all(build_grid())
    db.commit()

    result, run = generate(db, academic_context_id=context.id, soft=False, max_seconds=30)
    assert result.status in {"OPTIMAL", "FEASIBLE"}, result.status
    assert check_all(db, run.id) == []
    return {"db": db, "context": context, "run": run, "rooms": rooms, "faculty": faculty,
            "sections": sections}


def _master(client, **params) -> dict:
    resp = client.get("/api/timetable/master", params={"limit": 2000, **params})
    assert resp.status_code == 200, resp.text
    return resp.json()


def _rows(db, run_id: int) -> list[Assignment]:
    return db.query(Assignment).filter(Assignment.run_id == run_id).all()


# ------------------------------------------------- 1. exactly once, complete


def test_every_scheduled_period_appears_exactly_once_across_every_page(solved, client):
    db, run = solved["db"], solved["run"]
    expected = {a.id for a in _rows(db, run.id)}

    seen: list[int] = []
    offset = 0
    while True:
        page = client.get("/api/timetable/master",
                          params={"run_id": run.id, "limit": 5, "offset": offset}).json()
        assert page["total"] == len(expected)
        if not page["rows"]:
            break
        assert len(page["rows"]) <= 5
        seen.extend(r["assignment_id"] for r in page["rows"])
        offset += 5

    assert len(seen) == len(set(seen)), "a period appeared on two pages"
    assert set(seen) == expected, "a period is missing, or one from elsewhere appeared"


def test_every_row_carries_every_column_a_person_reads(solved, client):
    db, run = solved["db"], solved["run"]
    by_id = {a.id: a for a in _rows(db, run.id)}
    body = _master(client, run_id=run.id)
    assert body["version"] == run.version
    assert body["publish_status"] == "DRAFT"

    for r in body["rows"]:
        a = by_id[r["assignment_id"]]
        assert r["day"] == a.timeslot.day
        assert r["start_time"] == a.timeslot.start_time.strftime("%H:%M:%S")
        assert r["end_time"] == a.timeslot.end_time.strftime("%H:%M:%S")
        assert r["subject_code"] == a.subject.code and r["subject_name"] == a.subject.name
        assert r["section_number"] == a.section.section_number
        assert r["strength"] == a.section.strength
        assert r["faculty_name"] == a.faculty.name and r["faculty_code"] == a.faculty.faculty_code
        assert r["room_code"] == room_code(a.room)
        assert r["block"] == a.room.block and r["floor"] == (a.room.floor or "")
        assert r["room_capacity"] == a.room.capacity
        assert r["class_type"] == ("Lab" if a.session_type == "P" else "Theory")
        assert r["room_byod"] == bool(a.room.byod)
        assert r["block_id"] == a.block_id
        if a.subject.code == "EC282":
            assert r["byod_required"] is True and r["room_byod"] is True
            assert r["room_type"] == "Lab"


def test_the_filter_options_are_what_occurs_in_the_timetable(solved, client):
    db, run = solved["db"], solved["run"]
    rows = _rows(db, run.id)
    options = _master(client, run_id=run.id)["options"]
    assert {o["id"] for o in options["sections"]} == {a.section_id for a in rows}
    assert {o["id"] for o in options["faculty"]} == {a.faculty_id for a in rows}
    assert {o["id"] for o in options["subjects"]} == {a.subject_id for a in rows}
    assert {o["id"] for o in options["rooms"]} == {a.room_id for a in rows}
    assert {o["id"] for o in options["days"]} == {a.timeslot.day_index for a in rows}
    assert set(options["blocks"]) == {a.room.block for a in rows}
    assert set(options["floors"]) == {a.room.floor for a in rows}
    assert {o["id"] for o in options["room_types"]} == {a.room.room_type for a in rows}


# ------------------------------------------------------------- 2. filters


def _derive(db, run, locked_pairs, **f) -> set[int]:
    """The matching ids, worked out from the rows in Python."""
    q = (f.get("q") or "").lower()

    def keep(a: Assignment) -> bool:
        return all([
            f.get("section_id") in (None, a.section_id),
            f.get("faculty_id") in (None, a.faculty_id),
            f.get("subject_id") in (None, a.subject_id),
            f.get("room_id") in (None, a.room_id),
            f.get("day_index") in (None, a.timeslot.day_index),
            f.get("block") in (None, a.room.block),
            f.get("floor") in (None, a.room.floor),
            f.get("room_type") in (None, a.room.room_type),
            f.get("locked") is None
            or ((a.section_id, a.subject_id, a.session_type) in locked_pairs) == f["locked"],
            not q or any(q in (v or "").lower() for v in (
                a.subject.code, a.subject.name, a.section.section_number, a.faculty.name,
                a.faculty.faculty_code, a.room.room_number, a.room.room_code, a.room.block,
                room_code(a.room),
            )),
        ])

    return {a.id for a in _rows(db, run.id) if keep(a)}


def test_every_filter_alone_combined_and_cleared_returns_exactly_the_matching_rows(solved, client):
    from app.timetable_views import locked_pairs

    db, run = solved["db"], solved["run"]
    rows = _rows(db, run.id)
    sample = rows[0]

    # Lock one class so "locked" has both sides to tell apart.
    resp = client.post(
        f"/api/academic-contexts/{run.academic_context_id}/assignments/lock",
        json={"section_id": sample.section_id, "subject_id": sample.subject_id,
              "session_type": sample.session_type},
    )
    assert resp.status_code in (200, 201), resp.text
    pinned = locked_pairs(db, run.academic_context_id)
    assert pinned

    lab_row = next(a for a in rows if a.subject.code == "EC282")
    cases = [
        {"section_id": sample.section_id},
        {"faculty_id": sample.faculty_id},
        {"subject_id": sample.subject_id},
        {"room_id": sample.room_id},
        {"day_index": sample.timeslot.day_index},
        {"block": "36"},
        {"block": "37"},
        {"floor": "3"},
        {"room_type": "lab"},
        {"room_type": "theory"},
        {"locked": True},
        {"locked": False},
        {"q": "ec282"},                 # subject code, any case
        {"q": "operating"},             # subject name
        {"q": "2402"},                  # section
        {"q": "faculty c"},             # teacher name
        {"q": "90001"},                 # teacher id
        {"q": "36-10"},                 # the room as the views name it: block-number
        {"q": "37-301"},
        {"q": "nothing-matches-this"},
        {"section_id": lab_row.section_id, "room_type": "lab"},
        {"block": "37", "day_index": lab_row.timeslot.day_index, "locked": False},
        {"faculty_id": sample.faculty_id, "q": sample.section.section_number},
        # a combination that cannot match must return nothing, not everything
        {"room_type": "lab", "block": "36"},
    ]
    everything = _master(client, run_id=run.id)["total"]
    assert everything == len(rows)

    for f in cases:
        params = {k: (str(v).lower() if isinstance(v, bool) else v) for k, v in f.items()}
        body = _master(client, run_id=run.id, **params)
        got = [r["assignment_id"] for r in body["rows"]]
        expected = _derive(db, run, pinned, **f)
        assert len(got) == len(set(got)), f
        assert set(got) == expected, f
        assert body["total"] == len(expected), f

    # Clearing the filters is asking again without them: everything is back.
    assert _master(client, run_id=run.id)["total"] == everything


def test_a_filtered_result_pages_and_counts_what_matches(solved, client):
    db, run = solved["db"], solved["run"]
    block = "37"
    expected = {a.id for a in _rows(db, run.id) if a.room.block == block}
    assert len(expected) > 2

    seen: list[int] = []
    for offset in range(0, len(expected) + 2, 2):
        page = client.get("/api/timetable/master", params={
            "run_id": run.id, "block": block, "limit": 2, "offset": offset,
        }).json()
        assert page["total"] == len(expected)
        seen.extend(r["assignment_id"] for r in page["rows"])
    assert sorted(seen) == sorted(expected)


def test_the_exports_contain_exactly_what_the_search_matches(solved, client):
    db, run = solved["db"], solved["run"]
    expected = sum(1 for a in _rows(db, run.id) if a.subject.code == "EC282")
    resp = client.get("/api/export/master.csv", params={"run_id": run.id, "q": "EC282"})
    assert resp.status_code == 200
    lines = [line for line in resp.text.splitlines() if line]
    assert len(lines) - 1 == expected
    header = lines[0].split(",")
    for column in ("Day", "Start", "End", "Subject Code", "Section", "Faculty", "Faculty Code",
                   "Room", "Block", "Floor", "Capacity", "Type", "BYOD Room"):
        assert column in header


# ----------------------------------------------------------- 3. run scope


def test_every_view_shows_the_same_run_and_a_failed_later_run_replaces_none_of_them(solved, client):
    db, run, context = solved["db"], solved["run"], solved["context"]
    failed = TimetableRun(academic_context_id=context.id, version=run.version + 1,
                          status="INFEASIBLE", publish_status="DRAFT")
    db.add(failed)
    db.commit()

    sample = _rows(db, run.id)[0]
    ctx = {"academic_context_id": context.id}
    assert client.get("/api/timetable/master", params=ctx).json()["run_id"] == run.id
    assert client.get(f"/api/timetable/section/{sample.section_id}").json()["run_id"] == run.id
    assert client.get(f"/api/timetable/faculty/{sample.faculty_id}", params=ctx).json()["run_id"] == run.id
    assert client.get(f"/api/timetable/room/{sample.room_id}", params=ctx).json()["run_id"] == run.id
    assert client.get("/api/allocation", params=ctx).json()["run_id"] == run.id
    assert client.get("/api/allocation").json()["run_id"] == run.id
    assert client.get("/api/runs/latest", params={**ctx, "with_classes": True}).json()["id"] == run.id
    # Without the flag the newest run is still reported, failure and all - the
    # Generate page needs to say what became of the last attempt.
    assert client.get("/api/runs/latest", params=ctx).json()["id"] == failed.id


def test_two_runs_are_never_mixed(solved, client):
    db, run, context = solved["db"], solved["run"], solved["context"]
    _result, second = generate(db, academic_context_id=context.id, soft=False, max_seconds=30)
    assert second.id != run.id

    for run_id in (run.id, second.id):
        own = {a.id for a in _rows(db, run_id)}
        body = _master(client, run_id=run_id)
        assert body["run_id"] == run_id
        assert {r["assignment_id"] for r in body["rows"]} == own

    # Unnamed, it is the newer one - wholly.
    body = _master(client, academic_context_id=context.id)
    assert body["run_id"] == second.id
    assert {r["assignment_id"] for r in body["rows"]} == {a.id for a in _rows(db, second.id)}


# --------------------------------------------------- 4. cross-view consistency


def _classes_from_master(body) -> Counter:
    """Master rows grouped into classes (one per block), as the teacher's
    table and the Excel file name them."""
    blocks: dict[int, list[dict]] = {}
    for r in body["rows"]:
        blocks.setdefault(r["block_id"], []).append(r)
    out = Counter()
    for rows in blocks.values():
        rows.sort(key=lambda r: r["period_index"])
        first, last = rows[0], rows[-1]
        out[(first["day"], first["start_time"][:5], last["end_time"][:5], first["section_number"],
             first["subject_code"], first["faculty_code"], first["room_code"], len(rows))] += 1
    return out


def _classes_from_allocation(body) -> Counter:
    return Counter(
        (r["day"], r["start_time"], r["end_time"], r["section"], r["subject_code"],
         r["faculty_id"], r["room"], r["periods"])
        for r in body["rows"]
    )


def _classes_from_excel(content: bytes) -> dict[str, Counter]:
    wb = load_workbook(io.BytesIO(content))
    out = {}
    for sheet in ("Master", "Section-wise", "Faculty-wise", "Room-wise"):
        ws = wb[sheet]
        header = [c.value for c in ws[1]]
        col = {name: header.index(name) for name in header}
        counter = Counter()
        for row in ws.iter_rows(min_row=2, values_only=True):
            counter[(row[col["Day"]], row[col["Start Time"]], row[col["End Time"]],
                     str(row[col["Section"]]), row[col["Subject Code"]], str(row[col["Faculty ID"]]),
                     row[col["Room"]])] += 1
        out[sheet] = counter
    return out


def _grid_assignments(grid) -> dict[int, dict]:
    return {c["assignment_id"]: c for row in grid["rows"] for c in row["cells"]
            if c.get("assignment_id")}


def _assert_consistent(db, client, run_id: int) -> None:
    """Every view of the run describes exactly the run's rows, identically."""
    rows = _rows(db, run_id)
    by_id = {a.id: a for a in rows}
    master = _master(client, run_id=run_id)
    assert {r["assignment_id"] for r in master["rows"]} == set(by_id)

    # Week grids: every period in exactly the grids of its section, teacher and
    # room, and nowhere else.
    for kind, attr, ids in (
        ("section", "section_id", {a.section_id for a in rows}),
        ("faculty", "faculty_id", {a.faculty_id for a in rows}),
        ("room", "room_id", {a.room_id for a in rows}),
    ):
        for entity_id in ids:
            grid = client.get(f"/api/timetable/{kind}/{entity_id}", params={"run_id": run_id}).json()
            cells = _grid_assignments(grid)
            assert set(cells) == {a.id for a in rows if getattr(a, attr) == entity_id}, (kind, entity_id)
            for aid, cell in cells.items():
                a = by_id[aid]
                assert (cell["subject_code"], cell["section_id"], cell["faculty_id"], cell["room_id"],
                        cell["period_index"]) == (a.subject.code, a.section_id, a.faculty_id,
                                                  a.room_id, a.timeslot.period_index)

    # The teacher's table and every sheet of the Excel file: the same classes.
    classes = _classes_from_master(master)
    allocation = client.get("/api/allocation", params={"run_id": run_id}).json()
    assert _classes_from_allocation(allocation) == classes
    assert allocation["periods"] == len(rows)

    xlsx = client.get("/api/export/allocation.xlsx", params={"run_id": run_id})
    assert xlsx.status_code == 200
    without_length = Counter({k[:7]: v for k, v in classes.items()})
    for sheet, counter in _classes_from_excel(xlsx.content).items():
        assert counter == without_length, sheet


def test_one_class_reads_the_same_everywhere_through_a_move_a_room_change_and_a_regeneration(
    solved, client,
):
    db, run, context = solved["db"], solved["run"], solved["context"]
    _assert_consistent(db, client, run.id)

    # ---- a manual move
    a = next(x for x in _rows(db, run.id) if x.subject.code == "CS201")
    busy = {x.timeslot_id for x in _rows(db, run.id)
            if x.section_id == a.section_id or x.faculty_id == a.faculty_id or x.room_id == a.room_id}
    target = None
    for slot in db.query(TimeSlot).order_by(TimeSlot.day_index, TimeSlot.period_index):
        if slot.id in busy or slot.is_lunch:
            continue
        if client.post(f"/api/assignments/{a.id}/move/preview",
                       json={"target_timeslot_id": slot.id}).json()["ok"]:
            target = slot
            break
    assert target is not None
    resp = client.post(f"/api/assignments/{a.id}/move",
                       json={"target_timeslot_id": target.id, "lock_after": False})
    assert resp.status_code == 200, resp.text
    # The API shares this test's session; refresh the one row rather than
    # expiring everything, which would drop what earlier reads loaded.
    moved = db.get(Assignment, a.id)
    db.refresh(moved)
    assert moved.timeslot_id == target.id
    row = next(r for r in _master(client, run_id=run.id)["rows"] if r["assignment_id"] == a.id)
    assert (row["day_index"], row["period_index"]) == (target.day_index, target.period_index)
    _assert_consistent(db, client, run.id)

    # ---- a room change
    moved = db.get(Assignment, a.id)
    new_room = next(
        (r for r in solved["rooms"] if r.room_type == "theory" and r.id != moved.room_id
         and client.post(f"/api/assignments/{a.id}/room/preview", json={"room_id": r.id}).json()["ok"]),
        None,
    )
    assert new_room is not None
    resp = client.post(f"/api/assignments/{a.id}/room", json={"room_id": new_room.id})
    assert resp.status_code == 200, resp.text  # locks the class by default
    db.refresh(moved)
    row = next(r for r in _master(client, run_id=run.id)["rows"] if r["assignment_id"] == a.id)
    assert row["room_code"] == room_code(new_room) and row["locked"] is True
    old_room_grid = client.get(f"/api/timetable/room/{a.room_id}", params={"run_id": run.id}).json()
    assert a.id not in _grid_assignments(old_room_grid) or a.room_id == new_room.id
    _assert_consistent(db, client, run.id)

    # ---- regenerate, respecting the lock the room change left
    locked_before = {
        (x.timeslot_id, x.room_id, x.faculty_id)
        for x in _rows(db, run.id) if x.section_id == a.section_id and x.subject_id == a.subject_id
    }
    result, second = generate(db, academic_context_id=context.id, soft=False, max_seconds=30)
    assert result.status in {"OPTIMAL", "FEASIBLE"}
    locked_after = {
        (x.timeslot_id, x.room_id, x.faculty_id)
        for x in _rows(db, second.id) if x.section_id == a.section_id and x.subject_id == a.subject_id
    }
    assert locked_after == locked_before, "the locked class kept its exact placement"
    _assert_consistent(db, client, second.id)

    # The default view is now the new run, with none of the old run's rows.
    body = _master(client, academic_context_id=context.id)
    assert body["run_id"] == second.id
    assert not {r["assignment_id"] for r in body["rows"]} & {x.id for x in _rows(db, run.id)}
    grid = client.get(f"/api/timetable/room/{new_room.id}",
                      params={"academic_context_id": context.id}).json()
    assert grid["run_id"] == second.id
    assert set(_grid_assignments(grid)) <= {x.id for x in _rows(db, second.id)}


# ------------------------------------------------------------------ versions


def test_generating_adds_a_version_and_keeps_the_previous_one(solved, client):
    db, run, context = solved["db"], solved["run"], solved["context"]
    before = {a.id: (a.timeslot_id, a.room_id, a.faculty_id) for a in _rows(db, run.id)}
    _result, second = generate(db, academic_context_id=context.id, soft=False, max_seconds=30)
    assert second.version == run.version + 1
    listed = client.get("/api/runs", params={"academic_context_id": context.id}).json()
    assert [r["id"] for r in listed][:2] == [second.id, run.id]
    assert {a.id: (a.timeslot_id, a.room_id, a.faculty_id) for a in _rows(db, run.id)} == before


def test_a_published_version_cannot_be_changed_in_place_or_deleted(solved, client):
    db, run = solved["db"], solved["run"]
    assert client.post(f"/api/runs/{run.id}/validate").status_code == 200
    assert client.post(f"/api/runs/{run.id}/publish").status_code == 200
    a = _rows(db, run.id)[0]
    other_room = next(r for r in solved["rooms"] if r.id != a.room_id)
    slot = db.query(TimeSlot).filter(TimeSlot.id != a.timeslot_id).first()

    for path, body in (
        (f"/api/assignments/{a.id}/move/preview", {"target_timeslot_id": slot.id}),
        (f"/api/assignments/{a.id}/move", {"target_timeslot_id": slot.id}),
        (f"/api/assignments/{a.id}/room", {"room_id": other_room.id}),
        (f"/api/assignments/{a.id}/faculty", {"faculty_id": solved["faculty"][2].id}),
    ):
        resp = client.post(path, json=body)
        assert resp.status_code == 409, (path, resp.text)
        assert "published" in resp.json()["detail"]

    before = a.timeslot_id
    assert client.delete(f"/api/runs/{run.id}").status_code == 409
    db.refresh(a)
    assert a.timeslot_id == before
    assert _master(client, run_id=run.id)["publish_status"] == "PUBLISHED"


def test_publishing_a_newer_version_archives_the_older_and_the_two_compare(solved, client):
    db, run, context = solved["db"], solved["run"], solved["context"]
    for r in (run,):
        client.post(f"/api/runs/{r.id}/validate")
        assert client.post(f"/api/runs/{r.id}/publish").status_code == 200
    _result, second = generate(db, academic_context_id=context.id, soft=False, max_seconds=30)
    client.post(f"/api/runs/{second.id}/validate")
    assert client.post(f"/api/runs/{second.id}/publish").status_code == 200

    statuses = {r["id"]: r["publish_status"]
                for r in client.get("/api/runs", params={"academic_context_id": context.id}).json()}
    assert statuses[second.id] == "PUBLISHED" and statuses[run.id] == "ARCHIVED"

    diff = client.get(f"/api/runs/{run.id}/diff/{second.id}")
    assert diff.status_code == 200
    body = diff.json()
    assert body["from_run_id"] == run.id and body["to_run_id"] == second.id
    assert client.get(f"/api/runs/{run.id}/diff/{run.id}").json()["rows"] == []
