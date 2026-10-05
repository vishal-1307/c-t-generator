"""Starting clean, and the table the teacher actually reads.

Reset: every row of timetable data goes, the schema and the accounts stay, and
nothing is deleted without the exact confirmation phrase. The application has
to work straight afterwards - an empty database that cannot import is not a
reset, it is an outage.

Allocation: one row per class with the room it got and the facts that made the
room eligible, and an Excel file whose columns are the ones asked for.
"""
from __future__ import annotations

import io

import openpyxl

from app.allocation import EXCEL_COLUMNS
from app.models import (
    AcademicContext,
    AcademicContextRoom,
    Assignment,
    Faculty,
    Room,
    Section,
    Subject,
    TimeSlot,
    TimetableRun,
    User,
)
from app.reset import CONFIRMATION, data_counts, reset_application_data
from app.solver.run import generate
from tests.test_teacher_import import (
    INFRA_HEADERS,
    INFRA_ROWS,
    LOAD_HEADERS,
    LOAD_ROWS,
    _xlsx,
)

INFRA_CODES = {f"{r[0]}-{r[1]}" for r in INFRA_ROWS}


def _upload(client):
    return client.post(
        "/api/teacher/apply",
        files={
            "load": ("Load.xlsx", _xlsx(LOAD_HEADERS, LOAD_ROWS), "application/vnd.ms-excel"),
            "infra": ("Infra.xlsx", _xlsx(INFRA_HEADERS, INFRA_ROWS), "application/vnd.ms-excel"),
        },
    ).json()


def _generated(client, db):
    body = _upload(client)
    assert body["committed"] is True, body
    result, run = generate(db, body["academic_context_id"], max_seconds=30)
    assert result.ok, result.message
    return run


# ------------------------------------------------------------------ reset


def test_reset_empties_every_data_table_and_keeps_the_accounts(client, db_session):
    run = _generated(client, db_session)
    # The circular pair: a subject pinned to a room pinned to it.
    room = db_session.query(Room).first()
    subject = db_session.query(Subject).first()
    subject.fixed_room_id, room.fixed_subject_id = room.id, subject.id
    db_session.commit()
    users = db_session.query(User).count()
    assert run.id and users >= 1

    deleted = reset_application_data(db_session)
    assert deleted["room"] == 30 and deleted["assignment"] == 52

    db_session.expire_all()
    assert sum(data_counts(db_session).values()) == 0
    for model in (Room, Faculty, Subject, Section, AcademicContext, TimetableRun,
                  Assignment, TimeSlot, AcademicContextRoom):
        assert db_session.query(model).count() == 0, model.__name__
    assert db_session.query(User).count() == users


def test_the_application_works_straight_after_a_reset(client, db_session):
    _generated(client, db_session)
    reset_application_data(db_session)
    db_session.expire_all()

    run = _generated(client, db_session)
    rows = db_session.query(Assignment).filter_by(run_id=run.id).all()
    assert len(rows) == 52
    assert {db_session.get(Room, a.room_id).room_code for a in rows} <= INFRA_CODES


def test_the_reset_endpoint_needs_the_exact_phrase_and_an_admin(client, anon_client, db_session):
    _upload(client)
    before = sum(data_counts(db_session).values())

    preview = client.get("/api/admin/reset").json()
    assert preview["total_rows"] == before and preview["would_delete"]["room"] == 30
    assert sum(data_counts(db_session).values()) == before, "the preview deleted something"

    assert client.post("/api/admin/reset", json={"confirm": "yes"}).status_code == 400
    assert anon_client.post("/api/admin/reset", json={"confirm": CONFIRMATION}).status_code == 401
    assert sum(data_counts(db_session).values()) == before

    done = client.post("/api/admin/reset", json={"confirm": CONFIRMATION})
    assert done.status_code == 200, done.text
    assert done.json()["total_rows"] == before
    db_session.expire_all()
    assert sum(data_counts(db_session).values()) == 0


def test_the_reset_refuses_while_a_timetable_is_being_generated(client, monkeypatch):
    from app import jobs

    monkeypatch.setattr(jobs, "is_running", lambda: True)
    resp = client.post("/api/admin/reset", json={"confirm": CONFIRMATION})
    assert resp.status_code == 409


# ------------------------------------------------------------- allocation


def test_the_allocation_table_has_one_row_per_class_with_its_room(client, db_session):
    run = _generated(client, db_session)
    body = client.get("/api/allocation").json()     # no run named: the latest

    assert body["run_id"] == run.id
    assert body["periods"] == 52
    # Nine rows of the load sheet, four sessions a week each.
    assert body["classes"] == 36 == len(body["rows"])

    rooms = {r.room_code: r for r in db_session.query(Room).all()}
    for row in body["rows"]:
        assert row["room"] in INFRA_CODES
        assert row["room_capacity"] >= row["strength"]
        assert row["room_type"] == ("Lab" if row["class_type"] == "Lab" else "Classroom")
        if row["byod"]:
            assert row["room_byod"] is True
        assert rooms[row["room"]].capacity == row["room_capacity"]

    lab = next(r for r in body["rows"] if r["subject_code"] == "ECE102")
    assert lab["periods"] == 3, "a three-period lab is one row, not three"
    assert (lab["start_time"], lab["end_time"]) != ("", "")


def test_the_filters_narrow_the_rows_and_list_their_choices(client, db_session):
    _generated(client, db_session)
    body = client.get("/api/allocation").json()
    section = next(o for o in body["filters"]["sections"] if o["label"] == "2401")
    only = client.get(f"/api/allocation?section_id={section['id']}").json()["rows"]
    assert only and {r["section"] for r in only} == {"2401"}

    monday = client.get("/api/allocation?day=Monday").json()["rows"]
    assert {r["day"] for r in monday} <= {"Monday"}
    assert set(body["filters"]["days"]) <= {
        "Monday", "Tuesday", "Wednesday", "Thursday", "Friday"}


def test_the_excel_export_has_exactly_the_columns_asked_for(client, db_session):
    _generated(client, db_session)
    resp = client.get("/api/export/allocation.xlsx")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Master", "Section-wise", "Faculty-wise", "Room-wise"]
    # Every view is the same classes under the same header, ordered its own way.
    views = {name: list(wb[name].iter_rows(min_row=2, values_only=True)) for name in wb.sheetnames}
    assert all(sorted(rows, key=str) == sorted(views["Master"], key=str) for rows in views.values())
    for name in wb.sheetnames:
        assert wb[name].freeze_panes == "A2"
        assert wb[name].auto_filter.ref.startswith("A1:P")
    section_col = [c.value for c in wb["Section-wise"][1]].index("Section")
    in_section_order = [r[section_col] for r in views["Section-wise"]]
    assert in_section_order == sorted(in_section_order)

    ws = wb["Master"]
    header = [c.value for c in ws[1]]
    assert header == [name for name, _ in EXCEL_COLUMNS] == [
        "Day", "Start Time", "End Time", "Faculty ID", "Faculty Name",
        "Subject Code", "Subject Name", "Section", "Student Strength",
        "Class Type", "BYOD", "Room", "Block", "Floor", "Room Capacity", "Room Type",
    ]
    data = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(data) == 36
    col = {name: i for i, name in enumerate(header)}
    for row in data:
        assert f"{row[col['Block']]}-" in row[col["Room"]]
        assert row[col["Room"]] in INFRA_CODES
        assert row[col["BYOD"]] in ("Yes", "No")
        assert row[col["Room Capacity"]] >= row[col["Student Strength"]]


def test_there_is_nothing_to_show_before_anything_is_generated(client):
    assert client.get("/api/allocation").status_code == 404
    assert client.get("/api/export/allocation.xlsx").status_code == 404
