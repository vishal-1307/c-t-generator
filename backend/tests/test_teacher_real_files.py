"""The department's real Load.xlsx and Infra.xlsx, end to end.

The real files carry the names and IDs of real teachers, so they are never
committed. This test reads them from outside the repository:

* ``TEACHER_FILES_DIR`` - a folder holding ``Load.xlsx`` and ``Infra.xlsx``, or
* ``backend/tests/fixtures/private/`` - git-ignored, for a local copy.

Neither present, it skips - loudly, with the reason - rather than pass on
nothing. Present, it runs the whole two-file flow and checks the result from
the saved rows, not from what the solver says about itself:

    TEACHER_FILES_DIR=C:/path/to/files python -m pytest tests/test_teacher_real_files.py

Nothing here assumes a count particular to one version of the files. The
number of periods the timetable must hold is taken from the files themselves
(the upload's own summary), so an updated load sheet is tested as it is.
"""
from __future__ import annotations

import os
from collections import defaultdict
from pathlib import Path

import pytest

from app import jobs, timetable_checker
from app.models import Assignment, Room, Section, Subject, TimeSlot

HERE = Path(__file__).parent
CANDIDATES = [
    Path(os.environ["TEACHER_FILES_DIR"]) if os.environ.get("TEACHER_FILES_DIR") else None,
    HERE / "fixtures" / "private",
]


def _find() -> tuple[Path, Path] | None:
    for folder in CANDIDATES:
        if folder is None:
            continue
        load, infra = folder / "Load.xlsx", folder / "Infra.xlsx"
        if load.is_file() and infra.is_file():
            return load, infra
    return None


FILES = _find()
pytestmark = pytest.mark.skipif(
    FILES is None,
    reason="real teacher files not available - set TEACHER_FILES_DIR to a "
    "folder holding Load.xlsx and Infra.xlsx, or copy them into "
    "backend/tests/fixtures/private/",
)



@pytest.fixture
def workspace():
    """A client whose background solve thread shares the test database."""
    from fastapi.testclient import TestClient
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from sqlalchemy.pool import StaticPool

    from app.auth import hash_password
    from app.database import Base, get_db
    from app.main import app
    from app.models import User

    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    db = Session()
    db.add(User(username="admin", hashed_password=hash_password("pw12345678"),
                role="admin", is_active=True))
    db.commit()

    app.dependency_overrides[get_db] = lambda: db
    previous = jobs.session_factory
    jobs.use_session_factory(Session)
    client = TestClient(app)
    token = client.post(
        "/api/auth/login", json={"username": "admin", "password": "pw12345678"}
    ).json()["access_token"]
    client.headers["Authorization"] = f"Bearer {token}"

    yield {"client": client, "db": db}

    app.dependency_overrides.clear()
    jobs.use_session_factory(previous)
    db.close()


# How many times a week each class meets, when the file itself does not say.
#
# The copy of Load.xlsx this test was written against predates the Classes Per
# Week column, and the importer now refuses a load without it - nothing is
# guessed in production. So this test supplies the number the department's
# timetable was built on, 4, in memory, and only when the column is missing: a
# file that has it is uploaded exactly as it is. The original is never written.
LEGACY_CLASSES_PER_WEEK = 4


def _load_bytes(path: Path) -> bytes:
    import io

    import openpyxl

    from app.bulk_import.adapters import header_key
    from app.teacher_import.normalise import LOAD_ALIASES

    wb = openpyxl.load_workbook(path)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    if any(LOAD_ALIASES.get(header_key(str(h or ""))) == "sessions_per_week" for h in headers):
        return path.read_bytes()
    column = len(headers) + 1
    ws.cell(row=1, column=column, value="Classes Per Week")
    for r in range(2, ws.max_row + 1):
        if any(ws.cell(row=r, column=c).value not in (None, "") for c in range(1, column)):
            ws.cell(row=r, column=column, value=LEGACY_CLASSES_PER_WEEK)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(client, path: str):
    load, infra = FILES
    return client.post(
        path,
        files={
            "load": ("Load.xlsx", _load_bytes(load), "application/vnd.ms-excel"),
            "infra": ("Infra.xlsx", infra.read_bytes(), "application/vnd.ms-excel"),
        },
    )


def test_the_real_load_file_is_refused_until_it_says_how_often(workspace):
    """Uploaded exactly as it is, a load without Classes Per Week is refused,
    naming the column - rather than turned into a timetable for a guessed week."""
    load, infra = FILES
    from app.bulk_import.adapters import header_key
    from app.teacher_import.normalise import LOAD_ALIASES
    import openpyxl

    headers = [c.value for c in openpyxl.load_workbook(load).active[1]]
    if any(LOAD_ALIASES.get(header_key(str(h or ""))) == "sessions_per_week" for h in headers):
        pytest.skip("this copy of Load.xlsx already carries Classes Per Week")

    resp = workspace["client"].post(
        "/api/teacher/preview",
        files={
            "load": ("Load.xlsx", load.read_bytes(), "application/vnd.ms-excel"),
            "infra": ("Infra.xlsx", infra.read_bytes(), "application/vnd.ms-excel"),
        },
    )
    assert resp.status_code == 422
    problems = resp.json()["detail"]["problems"]
    assert len(problems) == 1 and "Classes Per Week" in problems[0]["message"]


def test_the_real_files_become_a_valid_timetable(workspace):
    client, db = workspace["client"], workspace["db"]

    # ------------------------------------------------ validate, then import
    preview = _upload(client, "/api/teacher/preview")
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["has_errors"] is False, [
        p for o in body["outcomes"] for p in o["problems"]]
    assert body["readiness"]["ready"] is True, body["readiness"]["blockers"]
    required = body["summary"]["required_periods"]
    assert required > 0

    applied = _upload(client, "/api/teacher/apply")
    assert applied.json()["committed"] is True
    context_id = applied.json()["academic_context_id"]

    # ------------------------------------------------------------ generate
    started = client.post(
        "/api/generate", json={"academic_context_id": context_id, "max_seconds": 60}
    )
    assert started.status_code == 202, started.text
    run_id = started.json()["id"]
    jobs.wait(run_id, timeout=300)
    db.expire_all()
    run = client.get(f"/api/runs/{run_id}").json()
    assert run["status"] in {"OPTIMAL", "FEASIBLE"}, run

    # --------------------------------------------------- check from the rows
    assert timetable_checker.check_all(db, run_id) == []
    assert client.get(f"/api/runs/{run_id}/check").json()["count"] == 0

    rows = db.query(Assignment).filter(Assignment.run_id == run_id).all()
    assert len(rows) == required, "every required period is placed, and no more"

    slots = {s.id: s for s in db.query(TimeSlot).all()}
    subjects = {s.id: s for s in db.query(Subject).all()}
    rooms = {r.id: r for r in db.query(Room).all()}
    sections = {s.id: s for s in db.query(Section).all()}

    assert {slots[a.timeslot_id].day for a in rows} <= {
        "Monday", "Tuesday", "Wednesday", "Thursday", "Friday"}
    for a in rows:
        assert rooms[a.room_id].capacity >= sections[a.section_id].strength
        if subjects[a.subject_id].byod_required:
            assert rooms[a.room_id].byod

    per_day = defaultdict(set)
    for a in rows:
        if subjects[a.subject_id].type != "practical":
            per_day[(a.section_id, a.subject_id, slots[a.timeslot_id].day)].add(a.block_id)
    assert all(len(v) == 1 for v in per_day.values())

    # Every room is one the real room list names - nothing from anywhere else.
    import io

    import openpyxl

    infra_rooms = set()
    listed = openpyxl.load_workbook(FILES[1], data_only=True).active
    for block, room, *_ in listed.iter_rows(min_row=2, values_only=True):
        if room is not None:
            infra_rooms.add(f"{block}-{room}")
    used = {rooms[a.room_id].room_code for a in rows}
    assert used <= infra_rooms, used - infra_rooms

    table = client.get(f"/api/allocation?run_id={run_id}").json()
    assert table["periods"] == required
    assert {r["room"] for r in table["rows"]} <= infra_rooms

    xlsx = client.get(f"/api/export/allocation.xlsx?run_id={run_id}")
    assert xlsx.status_code == 200
    sheet = openpyxl.load_workbook(io.BytesIO(xlsx.content)).active
    body = list(sheet.iter_rows(min_row=2, values_only=True))
    assert len(body) == table["classes"]
    assert {r[11] for r in body} <= infra_rooms       # the Room column

    print(f"\n   real files: {required} periods, status {run['status']}, "
          f"{run['solve_time_seconds']:.2f}s, 0 hard conflicts")
