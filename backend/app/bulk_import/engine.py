"""The shared import workflow: analyse, apply, templates, history.

Entity-agnostic by construction - it only ever talks to an `EntityAdapter`.
"""
from __future__ import annotations

import csv
import io
from typing import Any

from sqlalchemy.orm import Session

from ..models import ImportHistory
from .adapters import ADAPTERS, match_key, missing_required, normalise_row
from .parsing import parse_table
from .resolver import DbResolver, ReferenceResolver
from .types import (
    CREATE,
    DUPLICATE,
    INVALID,
    MODE_ADD_UPDATE,
    MODE_FULL_SYNC,
    MODES,
    UNCHANGED,
    UPDATE,
    Deactivation,
    ImportPreview,
    ParseError,
    RowResult,
)


def template_csv(entity: str) -> str:
    adapter = ADAPTERS[entity]
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=adapter.columns, lineterminator="\n")
    writer.writeheader()
    for sample in adapter.sample:
        writer.writerow({c: sample.get(c, "") for c in adapter.columns})
    return buf.getvalue()


def template_xlsx(entity: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    adapter = ADAPTERS[entity]
    wb = Workbook()
    ws = wb.active
    ws.title = adapter.label

    ws.append(adapter.columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for sample in adapter.sample:
        ws.append([sample.get(c, "") for c in adapter.columns])

    for i, column in enumerate(adapter.columns, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(14, len(column) + 4)

    # A second sheet documenting the columns, so the file is self-describing
    # once it has been mailed around and detached from this UI.
    notes = wb.create_sheet("Notes")
    notes.append(["Column", "Required", "Notes"])
    for cell in notes[1]:
        cell.font = Font(bold=True)
    for column in adapter.columns:
        notes.append([column, "yes" if column in adapter.required else "no", ""])
    notes.append([])
    notes.append(["About", adapter.description])
    notes.column_dimensions["A"].width = 18
    notes.column_dimensions["B"].width = 10
    notes.column_dimensions["C"].width = 90

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# -------------------------------------------------------------------- analysis


def analyse(
    db: Session,
    entity: str,
    data: bytes,
    filename: str,
    mode: str = MODE_ADD_UPDATE,
    resolver: ReferenceResolver | None = None,
) -> ImportPreview:
    """Parse + validate + match the whole file. Writes nothing.

    `resolver` decides what a row's references are allowed to point at.
    Left unset it resolves against the database, which is what a single-entity
    upload means. A whole-workbook import passes one that consults the rows
    staged from other sheets first, so a section-subject row can legitimately
    reference a section defined three sheets earlier and not yet written.
    """
    adapter = ADAPTERS[entity]
    if mode not in MODES:
        raise ParseError(f"mode must be one of {', '.join(MODES)}, got {mode!r}")

    headers, raw_rows = parse_table(data, filename)
    missing = missing_required(adapter, headers)
    if missing:
        raise ParseError(
            f"Missing required column(s): {', '.join(missing)}. "
            f"Expected header: {', '.join(adapter.columns)}"
        )

    existing = adapter.load_existing(db)
    # One resolver for the whole file: it prefetches each referenced table once
    # instead of the point query per row this used to run, which for a
    # five-thousand-row file meant fifteen thousand queries - and twice over,
    # because `apply` re-runs `analyse`.
    refs = resolver or DbResolver(db)
    seen: dict[tuple, int] = {}
    results: list[RowResult] = []

    for index, row in enumerate(raw_rows, start=2):  # header occupies row 1
        # Accept the column names the file actually uses before anything reads
        # it - see adapters.normalise_row.
        row = normalise_row(adapter, row)
        label = adapter.identity_label(row)
        payload, errors = adapter.validate(row, db, refs)
        identity = match_key(adapter, row, payload)

        if errors:
            results.append(
                RowResult(index, label, INVALID, "; ".join(errors),
                          payload=payload, identity_key=identity)
            )
            continue

        if identity in seen:
            results.append(
                RowResult(
                    index, label, DUPLICATE,
                    f"duplicate of row {seen[identity]} in this file (same identity)",
                    payload=payload, identity_key=identity,
                )
            )
            continue
        seen[identity] = index

        current = existing.get(identity)
        if current is None:
            results.append(
                RowResult(index, label, CREATE, "new record",
                          payload=payload, identity_key=identity)
            )
            continue

        changes = adapter.diff(current, payload)
        if not changes:
            results.append(
                RowResult(index, label, UNCHANGED, "identical to the stored record",
                          payload=payload, identity_key=identity)
            )
        else:
            results.append(
                RowResult(
                    index, label, UPDATE,
                    ", ".join(f"{c.field}: {c.old} → {c.new}" for c in changes),
                    changes=changes, payload=payload, identity_key=identity,
                )
            )

    deactivations: list[Deactivation] = []
    if mode == MODE_FULL_SYNC:
        for identity, obj in existing.items():
            if identity in seen:
                continue
            # Absent from the file. Reported for confirmation, never deleted -
            # and an already-inactive record is not re-reported as a change.
            if adapter.is_active(obj):
                deactivations.append(
                    Deactivation(
                        id=getattr(obj, "id", None), identity=adapter.record_label(obj),
                        currently_active=True, identity_key=identity,
                    )
                )

    return ImportPreview(
        entity=entity, mode=mode, filename=filename,
        columns=adapter.columns, rows=results, deactivations=deactivations,
    )


# ----------------------------------------------------------------------- apply


def apply(
    db: Session,
    entity: str,
    data: bytes,
    filename: str,
    mode: str = MODE_ADD_UPDATE,
    confirm_deactivations: bool = False,
    actor: str | None = None,
) -> ImportPreview:
    """Re-analyse, then apply in one transaction.

    Re-analysing rather than trusting a previously-returned preview means a
    row that became invalid since the preview (someone deleted the block,
    another import ran) is caught instead of blindly written.

    Refuses outright if any row is invalid: a reference-data import that is
    half-applied is worse than one that did not run, because nobody can tell
    which half landed.
    """
    adapter = ADAPTERS[entity]
    preview = analyse(db, entity, data, filename, mode)

    if preview.has_errors:
        _record_history(db, entity, filename, mode, preview, "rejected", actor,
                        "File contains invalid or duplicate rows; nothing was applied.")
        db.commit()
        preview.committed = False
        return preview

    if preview.deactivations and not confirm_deactivations:
        _record_history(db, entity, filename, mode, preview, "rejected", actor,
                        f"{len(preview.deactivations)} record(s) would be deactivated; "
                        "confirmation required.")
        db.commit()
        preview.committed = False
        return preview

    try:
        existing = adapter.load_existing(db)
        model = getattr(adapter, "model", None)

        if model is not None:
            # Bulk path (spec PART 21): one batched INSERT and one batched
            # UPDATE regardless of row count, instead of one query per row -
            # safe here because Room has no relationship columns to populate,
            # so every payload dict maps directly onto the table's columns.
            creates = [r.payload for r in preview.rows if r.verdict == CREATE]
            if creates:
                db.bulk_insert_mappings(model, creates)

            updates = [
                {**r.payload, "id": existing[r.identity_key].id}
                for r in preview.rows
                if r.verdict == UPDATE and r.identity_key in existing
            ]
            if updates:
                db.bulk_update_mappings(model, updates)

            if preview.deactivations:
                db.bulk_update_mappings(
                    model,
                    [{"id": d.id, "is_active": False} for d in preview.deactivations],
                )
        else:
            # Fallback for an adapter with no surrogate PK to bulk-write
            # against - junction/mapping tables (faculty<->subject,
            # section<->subject), which have a composite key and no `id`
            # column at all. One query per row; these tables are small
            # (bounded by faculty x subjects or sections x subjects, not by
            # raw row count the way rooms can be), so this is not the
            # bottleneck PART 21 is about.
            for result in preview.rows:
                if result.verdict == CREATE:
                    adapter.create(db, result.payload)
                elif result.verdict == UPDATE:
                    obj = existing.get(result.identity_key)
                    if obj is not None:
                        adapter.update(obj, result.payload)
            for deactivation in preview.deactivations:
                obj = existing.get(deactivation.identity_key)
                if obj is not None:
                    adapter.deactivate(db, obj)

        _record_history(db, entity, filename, mode, preview, "applied", actor)
        db.commit()
    except Exception:
        db.rollback()
        raise

    preview.committed = True
    return preview


def _record_history(
    db: Session,
    entity: str,
    filename: str,
    mode: str,
    preview: ImportPreview,
    status: str,
    actor: str | None,
    message: str | None = None,
) -> None:
    counts = preview.counts
    db.add(
        ImportHistory(
            entity=entity,
            filename=filename[:300],
            mode=mode,
            rows_processed=counts["total"],
            rows_created=counts[CREATE],
            rows_updated=counts[UPDATE],
            rows_unchanged=counts[UNCHANGED],
            rows_rejected=counts[INVALID] + counts[DUPLICATE],
            rows_deactivated=len(preview.deactivations) if status == "applied" else 0,
            status=status,
            message=message,
            actor=actor,
        )
    )
