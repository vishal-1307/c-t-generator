"""The generated timetable as the teacher reads it: one row per class.

Day, time, teacher, subject, section and size, and the room it was given with
the facts that made it eligible - block, floor, capacity, type, BYOD. One row
per class *session*, so a three-period lab is one row running 09:30-12:00
rather than three rows that differ only in the time.

The on-screen table and the Excel export both come from `allocation_rows`, so
the file cannot say anything the screen does not.
"""
from __future__ import annotations

from dataclasses import asdict, dataclass

from sqlalchemy.orm import Session

from .config import settings
from .domain import PRACTICAL
from .models import (
    AcademicContext,
    Assignment,
    Faculty,
    Room,
    Section,
    Subject,
    TimeSlot,
    TimetableRun,
)
# One rule for "which timetable" and one set of room-type words, for every view.
from .timetable_views import ROOM_TYPE_LABEL, SHOWN_STATUSES, current_run



@dataclass
class AllocationRow:
    day: str
    day_index: int
    start_time: str
    end_time: str
    periods: int
    faculty_id: str
    faculty_name: str
    subject_code: str
    subject_name: str
    section: str
    strength: int
    class_type: str
    byod: bool
    room: str
    block: str
    floor: str
    room_capacity: int
    room_type: str
    room_byod: bool
    # For filtering; not shown.
    faculty_pk: int
    section_pk: int
    room_pk: int

    def as_dict(self) -> dict:
        return asdict(self)


def latest_run(db: Session, academic_context_id: int | None = None) -> TimetableRun | None:
    """The timetable in force - ``timetable_views.current_run``, the one rule
    every view uses."""
    return current_run(db, academic_context_id)


def newest_run(db: Session, academic_context_id: int) -> TimetableRun | None:
    """A dataset's newest version with classes, published or not."""
    return (
        db.query(TimetableRun)
        .filter(TimetableRun.academic_context_id == academic_context_id,
                TimetableRun.status.in_(SHOWN_STATUSES))
        .order_by(TimetableRun.created_at.desc(), TimetableRun.id.desc())
        .first()
    )


def _hhmm(t) -> str:
    return t.strftime("%H:%M") if t is not None else ""


def allocation_rows(db: Session, run_id: int) -> list[AllocationRow]:
    """Every class session in the run, in day and time order."""
    rows = db.query(Assignment).filter(Assignment.run_id == run_id).all()
    if not rows:
        return []
    slots = {s.id: s for s in db.query(TimeSlot).all()}
    sections = {s.id: s for s in db.query(Section).all()}
    subjects = {s.id: s for s in db.query(Subject).all()}
    faculty = {f.id: f for f in db.query(Faculty).all()}
    rooms = {r.id: r for r in db.query(Room).all()}

    blocks: dict[int, list[Assignment]] = {}
    for a in rows:
        blocks.setdefault(a.block_id, []).append(a)

    out: list[AllocationRow] = []
    for block in blocks.values():
        block.sort(key=lambda a: (slots[a.timeslot_id].day_index,
                                  slots[a.timeslot_id].period_index))
        first, last = slots[block[0].timeslot_id], slots[block[-1].timeslot_id]
        a = block[0]
        section, subject = sections[a.section_id], subjects[a.subject_id]
        teacher, room = faculty[a.faculty_id], rooms[a.room_id]
        practical = (a.session_type or "L") == PRACTICAL
        out.append(AllocationRow(
            day=first.day,
            day_index=first.day_index,
            start_time=_hhmm(first.start_time),
            end_time=_hhmm(last.end_time),
            periods=len(block),
            faculty_id=teacher.faculty_code,
            faculty_name=teacher.name,
            subject_code=subject.code,
            subject_name=subject.name,
            section=section.section_number,
            strength=section.strength,
            class_type="Lab" if practical else "Theory",
            byod=bool(subject.byod_required)
            and (practical or settings.enforce_byod_on_lectures),
            room=room.room_code or (f"{room.block}-{room.room_number}"
                                    if room.block else room.room_number),
            block=room.block or "",
            floor=room.floor or "",
            room_capacity=room.capacity,
            room_type=ROOM_TYPE_LABEL.get(room.room_type, room.room_type),
            room_byod=bool(room.byod),
            faculty_pk=teacher.id,
            section_pk=section.id,
            room_pk=room.id,
        ))
    out.sort(key=lambda r: (r.day_index, r.start_time, r.section, r.subject_code))
    return out


def dataset_label(db: Session, run: TimetableRun) -> str:
    ctx = db.get(AcademicContext, run.academic_context_id)
    return ctx.label if ctx else ""


# Exactly the columns the teacher asked for, in that order.
EXCEL_COLUMNS: list[tuple[str, str]] = [
    ("Day", "day"),
    ("Start Time", "start_time"),
    ("End Time", "end_time"),
    ("Faculty ID", "faculty_id"),
    ("Faculty Name", "faculty_name"),
    ("Subject Code", "subject_code"),
    ("Subject Name", "subject_name"),
    ("Section", "section"),
    ("Student Strength", "strength"),
    ("Class Type", "class_type"),
    ("BYOD", "byod"),
    ("Room", "room"),
    ("Block", "block"),
    ("Floor", "floor"),
    ("Room Capacity", "room_capacity"),
    ("Room Type", "room_type"),
]


def allocation_workbook(rows: list[AllocationRow], title: str) -> bytes:
    """The timetable as four sheets of the same sixteen columns.

    Master is every class in the order of the week. The other three are the
    same rows grouped the way someone actually looks for their own part of it -
    a section, a teacher, a room - each still one table with one header and a
    filter, so nothing is lost by opening the "wrong" sheet and nothing needs a
    formula to read.
    """
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    from .exporters import _safe_cell

    def when(r: AllocationRow):
        return (r.day_index, r.start_time)

    views = [
        ("Master", sorted(rows, key=lambda r: (*when(r), r.section, r.subject_code))),
        ("Section-wise", sorted(rows, key=lambda r: (r.section, *when(r)))),
        ("Faculty-wise", sorted(rows, key=lambda r: (r.faculty_name.lower(), r.faculty_id, *when(r)))),
        ("Room-wise", sorted(rows, key=lambda r: (r.block, r.room, *when(r)))),
    ]

    wb = Workbook()
    wb.remove(wb.active)
    last = get_column_letter(len(EXCEL_COLUMNS))
    widths = [
        min(max(10, max([len(name)] + [len(str(getattr(r, key))) for r in rows]) + 2), 40)
        for name, key in EXCEL_COLUMNS
    ]
    for sheet_name, ordered in views:
        ws = wb.create_sheet(sheet_name)
        ws.append([name for name, _ in EXCEL_COLUMNS])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1D4ED8")
        for r in ordered:
            values = []
            for _, key in EXCEL_COLUMNS:
                v = getattr(r, key)
                if isinstance(v, bool):
                    v = "Yes" if v else "No"
                values.append(_safe_cell(v) if isinstance(v, str) else v)
            ws.append(values)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{last}{max(len(ordered) + 1, 1)}"
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
    wb.properties.title = title
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()